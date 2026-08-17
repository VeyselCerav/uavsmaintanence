from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.reports.datasets import ReportDocument
from apps.reports.labels import labels_for

FONT_CANDIDATES = [
    Path(__file__).resolve().parent / "fonts" / "DejaVuSans.ttf",
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
    Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
    Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
]

_FONT_NAME = "Helvetica"
_FONT_READY = False


def _register_font() -> str:
    global _FONT_NAME, _FONT_READY
    if _FONT_READY:
        return _FONT_NAME
    for path in FONT_CANDIDATES:
        if path.exists():
            pdfmetrics.registerFont(TTFont("ReportSans", str(path)))
            _FONT_NAME = "ReportSans"
            break
    _FONT_READY = True
    return _FONT_NAME


def render_xlsx(document: ReportDocument) -> bytes:
    workbook = Workbook()
    first = True
    for table in document.tables:
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        title = table.title[:31] or "Sheet"
        sheet.title = title
        sheet.append(table.headers)
        for row in table.rows:
            sheet.append(row)
        if not table.rows:
            sheet.append([labels_for(document.locale)["empty"]])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_pdf(document: ReportDocument) -> bytes:
    font = _register_font()
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    navy = colors.HexColor("#1B365D")
    muted = colors.HexColor("#5B6B7A")
    labels = labels_for(document.locale)
    title_style = ParagraphStyle("title", fontName=font, fontSize=14, leading=18, textColor=navy)
    meta_style = ParagraphStyle("meta", fontName=font, fontSize=9, leading=12, textColor=muted)
    heading_style = ParagraphStyle("h", fontName=font, fontSize=11, leading=14, textColor=navy)
    cell_style = ParagraphStyle("cell", fontName=font, fontSize=8, leading=10)
    story = [
        Paragraph(labels["app"], meta_style),
        Paragraph(document.title, title_style),
        Spacer(1, 4 * mm),
    ]
    if document.demo:
        story.append(Paragraph(labels["demo"], meta_style))
        story.append(Spacer(1, 3 * mm))
    for table in document.tables:
        story.append(Paragraph(table.title, heading_style))
        data = [[Paragraph(_as_html(cell), cell_style) for cell in table.headers]]
        if table.rows:
            for row in table.rows:
                data.append([Paragraph(_as_html(cell), cell_style) for cell in row])
        else:
            data.append([Paragraph(labels["empty"], cell_style)] + [""] * (len(table.headers) - 1))
        grid = Table(data, repeatRows=1)
        grid.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B365D")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D7DEE8")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("FONTNAME", (0, 0), (-1, -1), font),
                ]
            )
        )
        story.append(grid)
        story.append(Spacer(1, 5 * mm))
    doc.build(story)
    return buffer.getvalue()


def _as_html(value: str) -> str:
    return (
        str(value)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
